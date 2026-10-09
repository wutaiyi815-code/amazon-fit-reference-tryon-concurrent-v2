#!/usr/bin/env python3
"""Inventory, validate, and run Amazon apparel fit-reference try-on jobs.

Semantic image grouping and fit-reference choice intentionally live in a human-
reviewed selections.json. This program handles deterministic file/table routing,
Figure numbering, provider calls, logging, and resumable checkpoints.
"""

from __future__ import annotations

import argparse
import copy
import getpass
import io
import json
import os
import re
import shutil
import sys
import threading
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path
from typing import Any, Callable, Iterable

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

try:
    import requests
    from PIL import Image, ImageDraw, ImageFont, ImageOps
except ImportError as exc:
    raise SystemExit(
        "Missing runtime dependency. Use a Python environment containing requests and Pillow."
    ) from exc


CHAT_URL = "https://toapis.com/v1/chat/completions"
UPLOAD_URL = "https://toapis.com/v1/uploads/images"
TOAPIS_GENERATE_URL = "https://toapis.com/v1/images/generations"
RUNNINGHUB_GENERATE_URL = "https://www.runninghub.cn/openapi/v2/rhart-image-g-2-official/image-to-image"
RUNNINGHUB_QUERY_URL = "https://www.runninghub.cn/openapi/v2/query"
IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".webp"}
MODEL_VIEWS = {"front", "back", "side", "unknown"}
AREAS = {"upper", "lower", "full", "socks", "shoes", "accessory", "unknown"}
AREA_ORDER = {"upper": 0, "full": 1, "lower": 2, "socks": 3, "shoes": 4, "accessory": 5, "unknown": 6}
PRODUCT_TYPE_ORDER = {
    "tshirt": 0, "shirt": 1, "tank_top": 2, "hoodie": 3, "sweater": 4,
    "jacket": 5, "coat": 6, "vest": 7, "pants": 8, "shorts": 9,
    "skirt": 10, "dress": 11, "jumpsuit": 12, "socks": 13, "shoes": 14,
    "bag": 15, "hat": 16, "accessory": 17, "unknown": 18,
}
LOWER_HEM_TYPES = {
    "open",
    "elastic_cuff",
    "drawcord_released",
    "drawcord_cinched",
    "cropped",
    "unknown",
}
GENERATED_NAME_RE = re.compile(r"^(?:result_|contact_|figure_plan_|prompt_)", re.IGNORECASE)
MANDATORY_COMPANION_RE = re.compile(r"搭配品(?:平铺图)?", re.IGNORECASE)
PROMPT_SYSTEM_PATH = Path(__file__).resolve().parents[1] / "references" / "prompt_writer_system.md"


class WorkflowError(RuntimeError):
    pass


class WorkDirectoryLock:
    """Prevent two processes from mutating the same run checkpoint."""

    def __init__(self, path: Path):
        self.path = path
        self.handle: Any = None

    def __enter__(self) -> "WorkDirectoryLock":
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self.handle = self.path.open("a+b")
        try:
            if os.name == "nt":
                import msvcrt

                self.handle.seek(0, os.SEEK_END)
                if self.handle.tell() == 0:
                    self.handle.write(b"0")
                    self.handle.flush()
                self.handle.seek(0)
                msvcrt.locking(self.handle.fileno(), msvcrt.LK_NBLCK, 1)
            else:
                import fcntl

                fcntl.flock(self.handle.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
        except (OSError, IOError) as exc:
            self.handle.close()
            self.handle = None
            raise WorkflowError(
                f"运行目录已被另一个并发进程占用: {self.path.parent}"
            ) from exc
        return self

    def __exit__(self, exc_type: Any, exc: Any, traceback: Any) -> None:
        if self.handle is None:
            return
        try:
            self.handle.seek(0)
            if os.name == "nt":
                import msvcrt

                msvcrt.locking(self.handle.fileno(), msvcrt.LK_UNLCK, 1)
            else:
                import fcntl

                fcntl.flock(self.handle.fileno(), fcntl.LOCK_UN)
        finally:
            self.handle.close()
            self.handle = None


def log(message: str) -> None:
    stamp = datetime.now().strftime("%H:%M:%S")
    print(f"[{stamp}] {message}", flush=True)


def json_load(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8-sig"))


def json_save(path: Path, data: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_suffix(path.suffix + ".tmp")
    temp.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    for attempt in range(5):
        try:
            temp.replace(path)
            return
        except PermissionError:
            if attempt == 4:
                break
            time.sleep(0.2 * (attempt + 1))
    # Some Windows SMB shares deny atomic replace even though ordinary writes
    # are allowed. The per-run lock still guarantees a single checkpoint
    # writer, so preserve recoverability with a copy fallback.
    shutil.copyfile(temp, path)
    temp.unlink(missing_ok=True)


class CheckpointStore:
    """Serialize checkpoint writes while workers keep private task snapshots."""

    def __init__(self, path: Path, data: dict[str, Any]):
        self.path = path
        self.data = data
        self.lock = threading.RLock()

    def save(self) -> None:
        with self.lock:
            checkpoint_save(self.path, self.data)

    def get_task(self, key: str) -> dict[str, Any] | None:
        with self.lock:
            task = self.data.get("tasks", {}).get(key)
            return copy.deepcopy(task) if task is not None else None

    def set_task(self, key: str, task: dict[str, Any]) -> None:
        with self.lock:
            self.data.setdefault("tasks", {})[key] = copy.deepcopy(task)
            checkpoint_save(self.path, self.data)


def normalized_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if re.fullmatch(r"\d+\.0", text):
        text = text[:-2]
    return text


def normalized_key(value: Any) -> str:
    text = normalized_text(value).casefold()
    return re.sub(r"[\s_\-—–:：/／\\]+", "", text)


def normalized_sku(value: Any) -> str:
    return normalized_text(value).casefold()


def split_directory_tokens(name: str) -> list[str]:
    return [normalized_text(part) for part in re.split(r"[_/／]+", name) if normalized_text(part)]


def iter_images(directory: Path, recursive: bool = False) -> list[Path]:
    if not directory.is_dir():
        return []
    iterator = directory.rglob("*") if recursive else directory.iterdir()
    return sorted(
        [p.resolve() for p in iterator if p.is_file() and p.suffix.casefold() in IMAGE_EXTS],
        key=lambda p: str(p).casefold(),
    )


def root_model_images(sku_dir: Path) -> list[Path]:
    return [
        p for p in iter_images(sku_dir)
        if not GENERATED_NAME_RE.match(p.name)
        and not p.name.casefold().startswith("aigc_")
    ]


def resolve_library(requested: Path) -> tuple[Path, str | None]:
    requested = requested.expanduser()
    if requested.is_dir():
        return requested.resolve(), None
    raise WorkflowError(f"版型库不存在或不是文件夹: {requested}")


def match_tokenized_directories(parent: Path, requested_values: Iterable[str]) -> list[Path]:
    wanted = {normalized_key(value) for value in requested_values if normalized_text(value)}
    if not wanted or not parent.is_dir():
        return []
    matches: list[Path] = []
    for directory in sorted((p for p in parent.iterdir() if p.is_dir()), key=lambda p: p.name.casefold()):
        tokens = {normalized_key(token) for token in split_directory_tokens(directory.name)}
        if wanted.intersection(tokens):
            matches.append(directory.resolve())
    return matches


SKU_HEADER_ALIASES = tuple(
    normalized_key(x) for x in ("SKU", "SKU号", "产品SKU", "图片SKU", "商品SKU", "sku编码")
)
SPU_HEADER_ALIASES = tuple(
    normalized_key(x) for x in ("SPU(业务通用)", "SPU（业务通用）", "主SPU(amz)", "主SPU（amz）", "SPU")
) + SKU_HEADER_ALIASES


def find_header_row(
    worksheet: Any,
    required_headers: set[str],
    identity_aliases: Iterable[str] = SKU_HEADER_ALIASES,
) -> tuple[int, dict[str, int]]:
    for row_number, row in enumerate(worksheet.iter_rows(min_row=1, max_row=30, values_only=True), start=1):
        headers = {normalized_key(value): index for index, value in enumerate(row) if normalized_text(value)}
        has_identity = any(alias in headers for alias in identity_aliases)
        if has_identity and required_headers.intersection(headers):
            return row_number, headers
    raise WorkflowError(
        f"工作表 {worksheet.title!r} 的前 30 行中未找到 SKU 与所需字段表头"
    )


def find_identity_column(headers: dict[str, int], aliases: Iterable[str], label: str) -> int:
    for alias in aliases:
        if alias in headers:
            return headers[alias]
    raise WorkflowError(f"未找到 {label} 字段")


def read_excel_source(path: Path) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise WorkflowError("读取 Excel 需要 openpyxl") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    required_sheets = {"SPU数据", "AIGC需求"}
    missing = required_sheets.difference(workbook.sheetnames)
    if missing:
        raise WorkflowError(f"Excel 缺少工作表: {', '.join(sorted(missing))}")

    spu_sheet = workbook["SPU数据"]
    spu_header_row, spu_headers = find_header_row(
        spu_sheet,
        {normalized_key("二级类目-AM"), normalized_key("三级目录-AM")},
        SPU_HEADER_ALIASES,
    )
    spu_identity_col = find_identity_column(spu_headers, SPU_HEADER_ALIASES, "SPU")
    category2_col = spu_headers.get(normalized_key("二级类目-AM"))
    category3_col = spu_headers.get(normalized_key("三级目录-AM"))
    spu_data: list[dict[str, str]] = []
    for row in spu_sheet.iter_rows(min_row=spu_header_row + 1, values_only=True):
        spu = normalized_text(row[spu_identity_col] if spu_identity_col < len(row) else None)
        if not spu:
            continue
        spu_data.append({
            "spu": spu,
            "category2": normalized_text(row[category2_col] if category2_col is not None and category2_col < len(row) else None),
            "category3": normalized_text(row[category3_col] if category3_col is not None and category3_col < len(row) else None),
        })

    aigc_sheet = workbook["AIGC需求"]
    aigc_header_row, aigc_headers = find_header_row(aigc_sheet, {normalized_key("版型")})
    aigc_sku_col = find_identity_column(aigc_headers, SKU_HEADER_ALIASES, "SKU")
    aigc_spu_col = aigc_headers.get(normalized_key("SPU"))
    fit_col = aigc_headers[normalized_key("版型")]
    aigc_requirements: list[dict[str, str]] = []
    for row in aigc_sheet.iter_rows(min_row=aigc_header_row + 1, values_only=True):
        sku = normalized_text(row[aigc_sku_col] if aigc_sku_col < len(row) else None)
        if not sku:
            continue
        aigc_requirements.append({
            "sku": sku,
            "spu": normalized_text(row[aigc_spu_col] if aigc_spu_col is not None and aigc_spu_col < len(row) else None),
            "fit": normalized_text(row[fit_col] if fit_col < len(row) else None),
        })

    result = {
        "source": "excel",
        "source_path": str(path.resolve()),
        "spu_data": spu_data,
        "aigc_requirements": aigc_requirements,
    }
    workbook.close()
    return result


def read_table_json(path: Path) -> dict[str, Any]:
    data = json_load(path)
    if not isinstance(data.get("spu_data"), list) or not isinstance(data.get("aigc_requirements"), list):
        raise WorkflowError("表格 JSON 必须包含 spu_data 与 aigc_requirements 数组")
    return data


def unique_nonblank(records: list[dict[str, Any]], field: str, split_multi: bool = False) -> list[str]:
    result: list[str] = []
    seen: set[str] = set()
    for record in records:
        raw_value = normalized_text(record.get(field))
        values = re.split(r"[,，;；\n]+", raw_value) if split_multi else [raw_value]
        for raw_part in values:
            value = normalized_text(raw_part)
            key = normalized_key(value)
            if value and key not in seen:
                seen.add(key)
                result.append(value)
    return result


def table_info_for_sku(table_data: dict[str, Any], sku: str) -> dict[str, Any]:
    sku_key = normalized_sku(sku)
    aigc_rows = [row for row in table_data.get("aigc_requirements", []) if normalized_sku(row.get("sku")) == sku_key]
    spu_values = unique_nonblank(aigc_rows, "spu")
    spu_keys = {normalized_sku(value) for value in spu_values}
    if not spu_keys:
        # Backward-compatible normalized JSON may key SPU data directly by the full SKU.
        spu_keys = {sku_key}
    spu_rows = [
        row for row in table_data.get("spu_data", [])
        if normalized_sku(row.get("spu") or row.get("sku")) in spu_keys
    ]
    category2_values = unique_nonblank(spu_rows, "category2")
    category3_values = unique_nonblank(spu_rows, "category3")
    fit_values = unique_nonblank(aigc_rows, "fit", split_multi=True)
    conflicts = []
    for label, values in (
        ("二级类目-AM", category2_values),
        ("三级目录-AM", category3_values),
        ("版型", fit_values),
    ):
        if len(values) > 1:
            conflicts.append({"field": label, "values": values})
    return {
        "spu_row_count": len(spu_rows),
        "aigc_row_count": len(aigc_rows),
        "spu_values": spu_values,
        "category2_values": category2_values,
        "category3_values": category3_values,
        "fit_values": fit_values,
        "conflicts": conflicts,
    }


def library_candidates(library: Path, table_info: dict[str, Any]) -> dict[str, Any]:
    if table_info["conflicts"]:
        return {
            "category_match_level": None,
            "category_directories": [],
            "fit_directories": [],
            "images": [],
            "blocked_reason": "conflicting spreadsheet values",
        }
    category_directories = match_tokenized_directories(library, table_info["category3_values"])
    match_level = "三级目录-AM" if category_directories else None
    if not category_directories:
        category_directories = match_tokenized_directories(library, table_info["category2_values"])
        match_level = "二级类目-AM" if category_directories else None
    fit_directories: list[Path] = []
    for category_directory in category_directories:
        fit_directories.extend(match_tokenized_directories(category_directory, table_info["fit_values"]))
    unique_fit_dirs = list(dict.fromkeys(fit_directories))
    images: list[Path] = []
    for fit_directory in unique_fit_dirs:
        images.extend(iter_images(fit_directory, recursive=True))
    images = list(dict.fromkeys(images))
    blocked_reason = None
    if not table_info["fit_values"]:
        blocked_reason = "blank fit value"
    elif not category_directories:
        blocked_reason = "no exact category directory match"
    elif not unique_fit_dirs:
        blocked_reason = "no exact fit directory match"
    elif not images:
        blocked_reason = "matched fit directory has no images"
    return {
        "category_match_level": match_level,
        "category_directories": [str(p) for p in category_directories],
        "fit_directories": [str(p) for p in unique_fit_dirs],
        "images": [str(p) for p in images],
        "blocked_reason": blocked_reason,
    }


def load_font(size: int) -> ImageFont.ImageFont:
    candidates = [
        Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts" / "msyh.ttc",
        Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts" / "arial.ttf",
    ]
    for candidate in candidates:
        if candidate.is_file():
            try:
                return ImageFont.truetype(str(candidate), size=size)
            except Exception:
                pass
    return ImageFont.load_default()


def contact_sheets(images: list[Path], output_prefix: Path, title: str) -> list[str]:
    if not images:
        return []
    output_prefix.parent.mkdir(parents=True, exist_ok=True)
    columns, rows_per_page = 4, 4
    cell_w, cell_h, label_h, header_h = 320, 320, 54, 44
    per_page = columns * rows_per_page
    font = load_font(18)
    small_font = load_font(14)
    outputs: list[str] = []
    for page_index, offset in enumerate(range(0, len(images), per_page), start=1):
        page_images = images[offset:offset + per_page]
        canvas = Image.new("RGB", (columns * cell_w, header_h + rows_per_page * (cell_h + label_h)), "white")
        draw = ImageDraw.Draw(canvas)
        draw.text((12, 10), f"{title} — page {page_index}", fill="black", font=font)
        for local_index, path in enumerate(page_images):
            row, column = divmod(local_index, columns)
            x, y = column * cell_w, header_h + row * (cell_h + label_h)
            try:
                with Image.open(path) as raw:
                    image = ImageOps.exif_transpose(raw).convert("RGB")
                    image.thumbnail((cell_w - 16, cell_h - 16), Image.Resampling.LANCZOS)
                    px = x + (cell_w - image.width) // 2
                    py = y + (cell_h - image.height) // 2
                    canvas.paste(image, (px, py))
            except Exception as exc:
                draw.text((x + 8, y + 8), f"Unreadable: {exc}", fill="red", font=small_font)
            label = f"{offset + local_index + 1}. {path.name}"
            draw.rectangle((x, y + cell_h, x + cell_w, y + cell_h + label_h), fill="#f2f2f2")
            draw.text((x + 7, y + cell_h + 7), label[:48], fill="black", font=small_font)
        output = output_prefix.with_name(f"{output_prefix.name}_{page_index}.jpg")
        canvas.save(output, "JPEG", quality=88)
        outputs.append(str(output.resolve()))
    return outputs


def make_run_id() -> str:
    return datetime.now().strftime("%Y%m%d_%H%M%S")


def discover_sku_dirs(root: Path) -> list[Path]:
    if not root.is_dir():
        raise WorkflowError(f"目标根目录不存在: {root}")
    result = []
    for child in sorted((p for p in root.iterdir() if p.is_dir()), key=lambda p: p.name.casefold()):
        if child.name in {"AIGC_Result", "_aigc_fit_work"}:
            continue
        if (child / "上身").is_dir():
            result.append(child.resolve())
    if not result:
        raise WorkflowError("目标根目录的直接下级中未找到包含“上身”文件夹的 SKU")
    return result


def prepare(args: argparse.Namespace) -> int:
    root = Path(args.root).expanduser().resolve()
    library, library_note = resolve_library(Path(args.library))
    if library_note:
        log(library_note)
    else:
        log(f"Using fit library: {library}")
    if args.table_excel:
        table_path = Path(args.table_excel).expanduser().resolve()
        table_data = read_excel_source(table_path)
    else:
        table_path = Path(args.table_json).expanduser().resolve()
        table_data = read_table_json(table_path)

    run_id = make_run_id()
    work_dir = root / "_aigc_fit_work" / run_id
    work_dir.mkdir(parents=True, exist_ok=False)
    sku_entries: list[dict[str, Any]] = []
    selection_skus: list[dict[str, Any]] = []

    for sku_dir in discover_sku_dirs(root):
        sku = sku_dir.name
        models = root_model_images(sku_dir)
        # Match the source script: only top-level files in 上身 are products.
        # Nested folders such as 已处理 contain derived images and must be excluded.
        products = iter_images(sku_dir / "上身", recursive=False)
        local_upper = iter_images(sku_dir / "版型参考_上装", recursive=True)
        local_lower = iter_images(sku_dir / "版型参考_下装", recursive=True)
        table_info = table_info_for_sku(table_data, sku)
        lookup = library_candidates(library, table_info)
        sku_contact_dir = work_dir / "contacts" / sku
        contacts = {
            "models": contact_sheets(models, sku_contact_dir / "models", f"{sku} models"),
            "products": contact_sheets(products, sku_contact_dir / "products", f"{sku} products"),
            "local_upper": contact_sheets(local_upper, sku_contact_dir / "local_upper", f"{sku} local upper fit"),
            "local_lower": contact_sheets(local_lower, sku_contact_dir / "local_lower", f"{sku} local lower fit"),
            "library": contact_sheets([Path(p) for p in lookup["images"]], sku_contact_dir / "library", f"{sku} library fit"),
        }
        entry = {
            "sku": sku,
            "folder": str(sku_dir),
            "model_candidates": [str(p) for p in models],
            "product_candidates": [str(p) for p in products],
            "local_fit_candidates": {
                "upper": [str(p) for p in local_upper],
                "lower": [str(p) for p in local_lower],
            },
            "table": table_info,
            "library_lookup": lookup,
            "contact_sheets": contacts,
        }
        sku_entries.append(entry)
        selection_skus.append({
            "sku": sku,
            "models": [{"path": str(p), "view": ""} for p in models],
            "product_groups": [],
            "ignored_product_files": [],
        })
        log(
            f"Prepared {sku}: models={len(models)}, products={len(products)}, "
            f"local upper={len(local_upper)}, local lower={len(local_lower)}, library candidates={len(lookup['images'])}"
        )

    inventory = {
        "schema_version": 1,
        "run_id": run_id,
        "created_at": datetime.now().astimezone().isoformat(),
        "target_root": str(root),
        "work_dir": str(work_dir.resolve()),
        "table_source": table_data.get("source", "json"),
        "table_source_path": str(table_path),
        "library_requested": str(Path(args.library)),
        "library_resolved": str(library),
        "library_resolution_note": library_note,
        "skus": sku_entries,
    }
    selections = {
        "schema_version": 2,
        "run_id": run_id,
        "inventory": str((work_dir / "inventory.json").resolve()),
        "skus": selection_skus,
    }
    inventory_path = work_dir / "inventory.json"
    selections_path = work_dir / "selections.json"
    json_save(inventory_path, inventory)
    json_save(selections_path, selections)
    log(f"Inventory: {inventory_path}")
    log(f"Complete visual decisions in: {selections_path}")
    return 0


def path_key(value: str | Path) -> str:
    return os.path.normcase(str(Path(value).resolve()))


def validate_selection_data(inventory: dict[str, Any], selections: dict[str, Any]) -> list[str]:
    errors: list[str] = []
    selection_schema_version = selections.get("schema_version", 1)
    require_lower_hem = isinstance(selection_schema_version, int) and selection_schema_version >= 2
    if selections.get("run_id") != inventory.get("run_id"):
        errors.append("selections.run_id 与 inventory.run_id 不一致")
    inventory_by_sku = {entry["sku"]: entry for entry in inventory.get("skus", [])}
    selections_by_sku = {entry.get("sku"): entry for entry in selections.get("skus", [])}
    missing_skus = set(inventory_by_sku).difference(selections_by_sku)
    extra_skus = set(selections_by_sku).difference(inventory_by_sku)
    if missing_skus:
        errors.append(f"selections 缺少 SKU: {sorted(missing_skus)}")
    if extra_skus:
        errors.append(f"selections 含未知 SKU: {sorted(extra_skus)}")

    for sku, inv in inventory_by_sku.items():
        sel = selections_by_sku.get(sku)
        if not sel:
            continue
        expected_models = {path_key(p) for p in inv.get("model_candidates", [])}
        selected_models: set[str] = set()
        for model in sel.get("models", []):
            key = path_key(model.get("path", ""))
            selected_models.add(key)
            if key not in expected_models:
                errors.append(f"{sku}: 未知模特图 {model.get('path')}")
            if model.get("view") not in MODEL_VIEWS:
                errors.append(f"{sku}: 模特图 {model.get('path')} 的 view 无效")
        if selected_models != expected_models:
            errors.append(f"{sku}: 必须为每张根目录模特图填写且只填写一次 view")

        expected_products = {path_key(p) for p in inv.get("product_candidates", [])}
        used_products: set[str] = set()
        ignored_products: set[str] = set()
        for ignored in sel.get("ignored_product_files", []):
            if isinstance(ignored, str):
                errors.append(f"{sku}: ignored_product_files 必须包含 path 与 reason")
                continue
            key = path_key(ignored.get("path", ""))
            ignored_products.add(key)
            if key not in expected_products:
                errors.append(f"{sku}: 忽略了未知产品图 {ignored.get('path')}")
            if not normalized_text(ignored.get("reason")):
                errors.append(f"{sku}: 忽略产品图必须填写 reason")
            if MANDATORY_COMPANION_RE.search(Path(ignored.get("path", "")).name):
                errors.append(
                    f"{sku}: 搭配品属于必须上身的实际产品，不得忽略: {ignored.get('path')}"
                )

        group_ids: set[str] = set()
        area_counts: dict[str, int] = {}
        used_fit_refs: set[str] = set()
        local_upper = {path_key(p) for p in inv.get("local_fit_candidates", {}).get("upper", [])}
        local_lower = {path_key(p) for p in inv.get("local_fit_candidates", {}).get("lower", [])}
        library_refs = {path_key(p) for p in inv.get("library_lookup", {}).get("images", [])}
        for group in sel.get("product_groups", []):
            group_id = normalized_text(group.get("product_id"))
            if not group_id or group_id in group_ids:
                errors.append(f"{sku}: product_id 为空或重复: {group_id!r}")
            group_ids.add(group_id)
            area = group.get("garment_area")
            product_type = group.get("product_type", "unknown")
            if area not in AREAS:
                errors.append(f"{sku}/{group_id}: garment_area 无效")
                area = "unknown"
            area_counts[area] = area_counts.get(area, 0) + 1
            files = group.get("files") or []
            if not files:
                errors.append(f"{sku}/{group_id}: files 不能为空")
            for product_file in files:
                key = path_key(product_file.get("path", ""))
                if key not in expected_products:
                    errors.append(f"{sku}/{group_id}: 未知产品图 {product_file.get('path')}")
                if key in used_products:
                    errors.append(f"{sku}: 产品图被重复分组 {product_file.get('path')}")
                used_products.add(key)
                if product_file.get("view") not in MODEL_VIEWS:
                    errors.append(f"{sku}/{group_id}: 产品图 view 无效 {product_file.get('path')}")

            lower_hem = group.get("lower_hem")
            if product_type == "pants" and area == "lower":
                if require_lower_hem and not isinstance(lower_hem, dict):
                    errors.append(f"{sku}/{group_id}: schema_version 2 的长裤必须填写 lower_hem 人工判断")
                elif lower_hem is not None:
                    if not isinstance(lower_hem, dict):
                        errors.append(f"{sku}/{group_id}: lower_hem 必须是对象")
                    else:
                        lower_hem_type = normalized_text(lower_hem.get("type"))
                        if lower_hem_type not in LOWER_HEM_TYPES:
                            errors.append(
                                f"{sku}/{group_id}: lower_hem.type 无效，必须是 {sorted(LOWER_HEM_TYPES)} 之一"
                            )
                        if not normalized_text(lower_hem.get("reason")):
                            errors.append(f"{sku}/{group_id}: lower_hem 必须记录人工视觉判断 reason")
            elif lower_hem is not None:
                errors.append(f"{sku}/{group_id}: lower_hem 只适用于 garment_area=lower 的 pants")

            fit_reference = group.get("fit_reference")
            if area in {"full", "socks", "shoes", "accessory", "unknown"} and fit_reference:
                errors.append(f"{sku}/{group_id}: {area} 产品不得强行绑定上下装版型参考")
            local_for_area = local_upper if area == "upper" else local_lower if area == "lower" else set()
            if local_for_area and not fit_reference:
                errors.append(f"{sku}/{group_id}: 对应本地版型参考文件夹非空，必须选择最接近的一张")
            if fit_reference:
                ref_path = fit_reference.get("path", "")
                ref_key = path_key(ref_path)
                if ref_key in used_fit_refs:
                    errors.append(f"{sku}: 同一版型参考图被多个产品复用 {ref_path}")
                used_fit_refs.add(ref_key)
                if not normalized_text(fit_reference.get("reason")):
                    errors.append(f"{sku}/{group_id}: 版型参考必须记录视觉选择 reason")
                if local_for_area:
                    if ref_key not in local_for_area:
                        errors.append(f"{sku}/{group_id}: 本地参考非空时不能改用版型库")
                elif ref_key not in library_refs:
                    errors.append(f"{sku}/{group_id}: 版型参考不在允许的库候选中 {ref_path}")

        for area in ("upper", "lower", "full", "socks", "shoes"):
            if area_counts.get(area, 0) > 1:
                errors.append(f"{sku}: {area} 存在多个逻辑产品组，需人工明确取舍，不能自动叠穿")
        accounted = used_products.union(ignored_products)
        if accounted != expected_products:
            missing = sorted(expected_products.difference(accounted))
            extra = sorted(accounted.difference(expected_products))
            if missing:
                errors.append(f"{sku}: 以下产品图既未分组也未说明忽略: {missing}")
            if extra:
                errors.append(f"{sku}: 以下产品图重复或不属于候选: {extra}")
        overlap = used_products.intersection(ignored_products)
        if overlap:
            errors.append(f"{sku}: 产品图不能同时分组和忽略: {sorted(overlap)}")
        mandatory_companions = {
            path_key(path)
            for path in inv.get("product_candidates", [])
            if MANDATORY_COMPANION_RE.search(Path(path).name)
        }
        missing_companions = mandatory_companions.difference(used_products)
        if missing_companions:
            errors.append(
                f"{sku}: 以下搭配品必须进入 product_groups 并实际上身: {sorted(missing_companions)}"
            )
    return errors


def validate_command(args: argparse.Namespace) -> int:
    inventory = json_load(Path(args.inventory))
    selections = json_load(Path(args.selections))
    errors = validate_selection_data(inventory, selections)
    if errors:
        for error in errors:
            log(f"ERROR: {error}")
        log(f"Selection validation failed with {len(errors)} error(s)")
        return 2
    log("Selection validation passed")
    return 0


def choose_product_file(group: dict[str, Any], model_view: str) -> dict[str, Any]:
    files = group["files"]
    same_view = [item for item in files if item.get("view") == model_view]
    unknown = [item for item in files if item.get("view") == "unknown"]
    return (same_view or unknown or files)[0]


def sorted_groups(groups: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return sorted(groups, key=lambda group: (
        AREA_ORDER.get(group.get("garment_area"), 6),
        PRODUCT_TYPE_ORDER.get(group.get("product_type", "unknown"), 18),
        normalized_text(group.get("product_id")).casefold(),
    ))


def build_figure_plan(model: dict[str, Any], groups: list[dict[str, Any]]) -> list[dict[str, Any]]:
    plan: list[dict[str, Any]] = [{
        "figure": 1,
        "role": "model",
        "path": str(Path(model["path"]).resolve()),
        "view": model["view"],
    }]
    selected_groups: list[tuple[dict[str, Any], dict[str, Any], int]] = []
    for group in sorted_groups(groups):
        selected_file = choose_product_file(group, model["view"])
        figure = len(plan) + 1
        product_item = {
            "figure": figure,
            "role": "product",
            "path": str(Path(selected_file["path"]).resolve()),
            "product_id": group["product_id"],
            "product_type": group.get("product_type", "unknown"),
            "garment_area": group["garment_area"],
            "view": selected_file["view"],
            "has_fit_reference": bool(group.get("fit_reference")),
        }
        if group.get("lower_hem") is not None:
            product_item["lower_hem"] = copy.deepcopy(group["lower_hem"])
        plan.append(product_item)
        selected_groups.append((group, selected_file, figure))
    for group, _selected_file, product_figure in selected_groups:
        fit_reference = group.get("fit_reference")
        if not fit_reference:
            continue
        plan.append({
            "figure": len(plan) + 1,
            "role": "fit_reference",
            "path": str(Path(fit_reference["path"]).resolve()),
            "for_product_id": group["product_id"],
            "for_product_figure": product_figure,
            "garment_area": group["garment_area"],
            "reason": fit_reference["reason"],
        })
    return plan


def validate_figure_product_coverage(
    sku: str,
    groups: list[dict[str, Any]],
    plan: list[dict[str, Any]],
) -> list[str]:
    """Independent pre-upload gate: every logical product must be sent once."""
    expected = [normalized_text(group.get("product_id")) for group in groups]
    actual = [
        normalized_text(item.get("product_id"))
        for item in plan
        if item.get("role") == "product"
    ]
    errors: list[str] = []
    if len(actual) != len(expected):
        errors.append(
            f"{sku}: Figure 产品数 {len(actual)} 与逻辑产品组数 {len(expected)} 不一致"
        )
    missing = sorted(set(expected).difference(actual))
    duplicated = sorted(product_id for product_id in set(actual) if actual.count(product_id) > 1)
    extra = sorted(set(actual).difference(expected))
    if missing:
        errors.append(f"{sku}: Figure 计划缺少产品组 {missing}")
    if duplicated:
        errors.append(f"{sku}: Figure 计划重复产品组 {duplicated}")
    if extra:
        errors.append(f"{sku}: Figure 计划包含未知产品组 {extra}")
    return errors


def lower_hem_contract(product: dict[str, Any]) -> str | None:
    lower_hem = product.get("lower_hem")
    if not isinstance(lower_hem, dict):
        return None
    figure = product["figure"]
    lower_hem_type = normalized_text(lower_hem.get("type"))
    if lower_hem_type == "open":
        return (
            f"Keep the Figure {figure} pants full-length with open, non-elastic hems. "
            "Let each hem fall naturally onto and partially over the shoe upper with gravity-driven folds. "
            "Do not tuck, cuff, cinch, gather, or bunch the hems around the ankles."
        )
    if lower_hem_type == "drawcord_released":
        return (
            f"Keep the Figure {figure} pants' adjustable hem drawcords fully released and the leg openings loose. "
            "Let the hems fall naturally onto and partially over the shoe uppers; "
            "do not cinch, cuff, tuck, gather, or bunch them around the ankles."
        )
    return None


def lower_hem_contracts(plan: list[dict[str, Any]]) -> list[str]:
    return [
        contract
        for product in plan
        if product.get("role") == "product"
        for contract in [lower_hem_contract(product)]
        if contract
    ]


def apply_lower_hem_contracts(prompt: str, plan: list[dict[str, Any]]) -> str:
    contracts = lower_hem_contracts(plan)
    missing = [contract for contract in contracts if contract not in prompt]
    if not missing:
        return prompt.strip()
    return f"{prompt.strip()}\n\n" + "\n".join(missing)


def validate_prompt_contracts(prompt: str, plan: list[dict[str, Any]]) -> list[str]:
    return [
        f"最终提示词缺少确定性裤脚约束: {contract}"
        for contract in lower_hem_contracts(plan)
        if contract not in prompt
    ]


def build_mapping_text(plan: list[dict[str, Any]], extra_requirements: str = "") -> str:
    lines = [
        "AUTHORITATIVE FIGURE ROLE AND CORRESPONDENCE MAP:",
        "Figure 1 is the target model/base image.",
        "All product Figures are actual products to put on the model. All fit-reference Figures appear only after all products.",
    ]
    products = [item for item in plan if item["role"] == "product"]
    references = [item for item in plan if item["role"] == "fit_reference"]
    ref_by_product = {item["for_product_id"]: item for item in references}
    for product in products:
        ref = ref_by_product.get(product["product_id"])
        lines.append(
            f"Figure {product['figure']} is the actual {product['garment_area']} product "
            f"(product_id={product['product_id']}, product_type={product['product_type']}, filename={Path(product['path']).name})."
        )
        contract = lower_hem_contract(product)
        if contract:
            lines.append(
                "AUTHORITATIVE LOWER-HEM CONTRACT FOR THIS PRODUCT: " + contract
            )
        if ref:
            lines.append(
                f"Figure {ref['figure']} is the fit-and-silhouette reference ONLY for the actual product in "
                f"Figure {product['figure']}. It must not affect any other product and must not supply color, fabric, pattern, logo, hardware, or design details."
            )
        else:
            lines.append(
                f"Figure {product['figure']} has NO fit-reference Figure. Do not assign any other reference to it."
            )
    if extra_requirements:
        lines.extend(["", "ADDITIONAL USER REQUIREMENTS FROM THIS SKU'S prompt.txt:", extra_requirements])
    return "\n".join(lines)


def normalize_api_key(value: str) -> str:
    value = value.strip()
    if value and not value.startswith("Bearer "):
        return f"Bearer {value}"
    return value


def process_image_for_upload(path: Path, max_size: int = 2048) -> io.BytesIO:
    with Image.open(path) as original:
        image = ImageOps.exif_transpose(original)
        image.load()
        if max(image.size) > max_size:
            scale = max_size / max(image.size)
            image = image.resize((max(1, int(image.width * scale)), max(1, int(image.height * scale))), Image.Resampling.LANCZOS)
        suffix = path.suffix.casefold()
        stream = io.BytesIO()
        if suffix in {".jpg", ".jpeg"}:
            image = image.convert("RGB")
            image.save(stream, "JPEG", quality=90)
            stream.mime_type = "image/jpeg"  # type: ignore[attr-defined]
        else:
            if image.mode not in {"RGB", "RGBA", "L", "LA"}:
                image = image.convert("RGBA" if "A" in image.getbands() else "RGB")
            image.save(stream, "PNG")
            stream.mime_type = "image/png"  # type: ignore[attr-defined]
        stream.seek(0)
        return stream


def upload_image(api_key: str, path: Path) -> str:
    retry_codes = {429, 500, 502, 503, 504, 520, 521, 522, 523, 524}
    for attempt in range(5):
        if attempt:
            time.sleep(min(5 * (2 ** (attempt - 1)), 30))
        stream = process_image_for_upload(path)
        files = {"file": (path.name, stream, getattr(stream, "mime_type", "image/png"))}
        try:
            response = requests.post(UPLOAD_URL, headers={"Authorization": api_key}, files=files, timeout=120)
            if response.status_code == 200:
                body = response.json()
                url = body.get("url") or body.get("data", {}).get("url")
                if url:
                    return url
            if response.status_code not in retry_codes:
                raise WorkflowError(f"上传失败 {path.name}: HTTP {response.status_code} {response.text[:500]}")
            log(f"Upload retryable failure {path.name}: HTTP {response.status_code}")
        except requests.RequestException as exc:
            if attempt == 4:
                raise WorkflowError(f"上传网络失败 {path.name}: {exc}") from exc
    raise WorkflowError(f"上传重试耗尽: {path.name}")


def extract_message_text(body: dict[str, Any]) -> str:
    choice = (body.get("choices") or [{}])[0]
    content = (choice.get("message") or {}).get("content", "")
    if isinstance(content, str):
        return content.strip()
    if isinstance(content, list):
        return "\n".join(str(item.get("text", "")) for item in content if isinstance(item, dict)).strip()
    return ""


def generate_prompt(api_key: str, language_model: str, image_urls: list[str], mapping_text: str) -> str:
    system_prompt = PROMPT_SYSTEM_PATH.read_text(encoding="utf-8-sig").strip()
    content: list[dict[str, Any]] = [{"type": "text", "text": mapping_text}]
    for index, url in enumerate(image_urls, start=1):
        content.append({"type": "text", "text": f"Figure {index}:"})
        content.append({"type": "image_url", "image_url": {"url": url}})
    payload = {
        "model": language_model,
        "messages": [
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": content},
        ],
        "max_tokens": 1600,
    }
    for attempt in range(3):
        if attempt:
            time.sleep(5 * attempt)
        try:
            response = requests.post(CHAT_URL, headers={"Authorization": api_key, "Content-Type": "application/json"}, json=payload, timeout=120)
            if response.status_code == 200:
                prompt = extract_message_text(response.json())
                if prompt:
                    return prompt
            log(f"Prompt writer attempt {attempt + 1} failed: HTTP {response.status_code} {response.text[:500]}")
        except requests.RequestException as exc:
            log(f"Prompt writer attempt {attempt + 1} network error: {exc}")
    raise WorkflowError("提示词生成失败；不使用泛化回退提示词继续付费生图")


def checkpoint_save(path: Path, checkpoint: dict[str, Any]) -> None:
    checkpoint["updated_at"] = datetime.now().astimezone().isoformat()
    json_save(path, checkpoint)


def extract_result_url(body: dict[str, Any]) -> str | None:
    """Read an output URL from the response shapes used by supported providers."""

    def walk(value: Any) -> str | None:
        if isinstance(value, dict):
            for key in ("url", "image_url", "imageUrl"):
                candidate = value.get(key)
                if isinstance(candidate, str) and candidate.strip():
                    return candidate.strip()
            for key in ("data", "items", "images", "output", "result"):
                if key in value:
                    candidate = walk(value[key])
                    if candidate:
                        return candidate
        elif isinstance(value, list):
            for item in value:
                candidate = walk(item)
                if candidate:
                    return candidate
        return None

    for key in ("result", "output", "data"):
        if key in body:
            candidate = walk(body[key])
            if candidate:
                return candidate
    return None


def poll_toapis(
    api_key: str,
    task_id: str,
    task: dict[str, Any],
    persist_task: Callable[[dict[str, Any]], None],
) -> bytes:
    poll_url = f"{TOAPIS_GENERATE_URL}/{task_id}"
    started = time.monotonic()
    last_heartbeat = 0.0
    while time.monotonic() - started < 900:
        try:
            response = requests.get(poll_url, headers={"Authorization": api_key}, timeout=60)
            if response.status_code == 200:
                body = response.json()
                status = str(body.get("status") or "unknown").casefold()
                task["provider_status"] = status
                persist_task(task)
                if status in {"completed", "success", "succeeded"}:
                    result_url = extract_result_url(body)
                    if not result_url:
                        raise WorkflowError(f"ToAPIs task {task_id} completed without result URL")
                    task["result_url"] = result_url
                    persist_task(task)
                    return requests.get(result_url, timeout=180).content
                if status in {"failed", "error", "cancelled", "canceled"}:
                    raise WorkflowError(f"ToAPIs task {task_id} failed: {str(body)[:1000]}")
            else:
                log(f"ToAPIs poll HTTP {response.status_code}: {response.text[:300]}")
        except requests.RequestException as exc:
            log(f"ToAPIs poll network check failed: {exc}")
        elapsed = time.monotonic() - started
        if elapsed - last_heartbeat >= 40:
            log(f"ToAPIs task {task_id} still alive; last status={task.get('provider_status', 'unknown')}; next check in 8s")
            last_heartbeat = elapsed
        time.sleep(8)
    raise WorkflowError(f"ToAPIs task {task_id} exceeded 15-minute bounded polling window")


def submit_toapis(api_key: str, model: str, prompt: str, image_urls: list[str], aspect_ratio: str, resolution: str) -> str:
    payload: dict[str, Any] = {
        "model": model,
        "prompt": prompt,
        "size": aspect_ratio,
        "n": 1,
        "image_urls": image_urls,
        "metadata": {"resolution": resolution},
    }
    if model == "gpt-image-2":
        payload.update({"resolution": resolution.casefold(), "quality": "medium", "response_format": "url"})
    for attempt in range(5):
        if attempt:
            time.sleep(min(5 * (2 ** (attempt - 1)), 30))
        try:
            response = requests.post(
                TOAPIS_GENERATE_URL,
                headers={"Authorization": api_key, "Content-Type": "application/json"},
                json=payload,
                timeout=60,
            )
            if response.status_code == 200:
                body = response.json()
                task_id = body.get("id") or body.get("task_id")
                if task_id:
                    return str(task_id)
            log(f"ToAPIs submit attempt {attempt + 1} failed: HTTP {response.status_code} {response.text[:500]}")
        except requests.RequestException as exc:
            log(f"ToAPIs submit attempt {attempt + 1} network error: {exc}")
    raise WorkflowError("ToAPIs submit retries exhausted")


def poll_runninghub(
    api_key: str,
    task_id: str,
    task: dict[str, Any],
    persist_task: Callable[[dict[str, Any]], None],
) -> bytes:
    started = time.monotonic()
    last_heartbeat = 0.0
    while time.monotonic() - started < 900:
        try:
            response = requests.post(
                RUNNINGHUB_QUERY_URL,
                headers={"Authorization": api_key, "Content-Type": "application/json"},
                json={"taskId": task_id},
                timeout=60,
            )
            if response.status_code == 200:
                body = response.json()
                result_body = body.get("data") or body
                status = str(result_body.get("status") or "unknown").upper()
                task["provider_status"] = status
                persist_task(task)
                if status == "SUCCESS":
                    results = result_body.get("results") or []
                    first = results[0] if results else {}
                    result_url = first.get("url") if isinstance(first, dict) else None
                    if not result_url:
                        raise WorkflowError(f"RunningHub task {task_id} completed without result URL")
                    task["result_url"] = result_url
                    persist_task(task)
                    return requests.get(result_url, timeout=180).content
                if status == "FAILED":
                    raise WorkflowError(f"RunningHub task {task_id} failed: {str(result_body)[:1000]}")
            else:
                log(f"RunningHub poll HTTP {response.status_code}: {response.text[:300]}")
        except requests.RequestException as exc:
            log(f"RunningHub poll network check failed: {exc}")
        elapsed = time.monotonic() - started
        if elapsed - last_heartbeat >= 40:
            log(f"RunningHub task {task_id} still alive; last status={task.get('provider_status', 'unknown')}; next check in 8s")
            last_heartbeat = elapsed
        time.sleep(8)
    raise WorkflowError(f"RunningHub task {task_id} exceeded 15-minute bounded polling window")


def submit_runninghub(api_key: str, prompt: str, image_urls: list[str], aspect_ratio: str, resolution: str) -> str:
    if len(image_urls) > 10:
        raise WorkflowError("RunningHub image-to-image 最多允许 10 张输入图，当前 Figure 计划超过限制")
    payload = {
        "prompt": prompt,
        "imageUrls": image_urls,
        "aspectRatio": aspect_ratio if aspect_ratio in {"auto", "1:1", "2:3", "3:2", "3:4", "4:3", "4:5", "5:4", "9:16", "16:9", "21:9"} else "3:4",
        "resolution": resolution if resolution in {"1K", "2K", "4K"} else "4K",
        "quality": "medium",
    }
    for attempt in range(3):
        if attempt:
            time.sleep(8 * attempt)
        try:
            response = requests.post(
                RUNNINGHUB_GENERATE_URL,
                headers={"Authorization": api_key, "Content-Type": "application/json"},
                json=payload,
                timeout=120,
            )
            if response.status_code == 200:
                body = response.json()
                task_id = body.get("taskId") or (body.get("data") or {}).get("taskId")
                if task_id:
                    return str(task_id)
            log(f"RunningHub submit attempt {attempt + 1} failed: HTTP {response.status_code} {response.text[:500]}")
        except requests.RequestException as exc:
            log(f"RunningHub submit attempt {attempt + 1} network error: {exc}")
    raise WorkflowError("RunningHub submit retries exhausted")


def prompt_secret(label: str, env_name: str) -> str:
    value = os.environ.get(env_name, "").strip()
    if not value:
        value = getpass.getpass(f"{label}: ").strip()
    if not value:
        raise WorkflowError(f"缺少 {label}")
    return normalize_api_key(value)


def read_sku_prompt(sku_dir: Path) -> str:
    prompt_path = sku_dir / "prompt.txt"
    if not prompt_path.is_file():
        return ""
    return prompt_path.read_text(encoding="utf-8-sig").strip()


def task_key(sku: str, model_path: str) -> str:
    return f"{sku}::{path_key(model_path)}"


def atomic_write_bytes(path: Path, content: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_name(f"{path.name}.{threading.get_ident()}.tmp")
    try:
        temp.write_bytes(content)
        temp.replace(path)
    finally:
        if temp.exists():
            temp.unlink()


def worker_count(value: str) -> int:
    try:
        count = int(value)
    except ValueError as exc:
        raise argparse.ArgumentTypeError("--max-workers 必须是整数") from exc
    if not 1 <= count <= 8:
        raise argparse.ArgumentTypeError("--max-workers 必须在 1 到 8 之间")
    return count


def execute_generation_job(
    job: dict[str, Any],
    store: CheckpointStore,
    toapis_key: str,
    runninghub_key: str,
    generation_model: str,
    aspect_ratio: str,
    resolution: str,
) -> Path:
    key = job["key"]
    output_path = Path(job["output_path"])
    task = store.get_task(key)
    if task is None:
        raise WorkflowError(f"checkpoint 缺少任务: {key}")

    def persist(current: dict[str, Any]) -> None:
        store.set_task(key, current)

    try:
        provider_task_id = task.get("provider_task_id")
        if provider_task_id:
            log(f"Recovering existing task {provider_task_id} for {task['sku']}/{Path(task['model_path']).name}")
        else:
            task["status"] = "submitting"
            task.pop("error", None)
            persist(task)
            log(f"Submitting paid generation for {task['sku']}/{Path(task['model_path']).name}")
            if generation_model == "gpt-image-2-rh":
                provider_task_id = submit_runninghub(
                    runninghub_key,
                    task["prompt"],
                    task["ordered_image_urls"],
                    aspect_ratio,
                    resolution,
                )
            else:
                provider_task_id = submit_toapis(
                    toapis_key,
                    generation_model,
                    task["prompt"],
                    task["ordered_image_urls"],
                    aspect_ratio,
                    resolution,
                )
            task["provider_task_id"] = provider_task_id
            task["status"] = "submitted"
            persist(task)
            log(f"Persisted provider task ID: {provider_task_id} ({task['sku']})")

        task["status"] = "polling"
        task.pop("error", None)
        persist(task)
        if task["provider"] == "runninghub":
            image_bytes = poll_runninghub(
                runninghub_key,
                str(provider_task_id),
                task,
                persist,
            )
        else:
            image_bytes = poll_toapis(
                toapis_key,
                str(provider_task_id),
                task,
                persist,
            )
        atomic_write_bytes(output_path, image_bytes)
        task["status"] = "completed"
        task["output_path"] = str(output_path)
        task.pop("error", None)
        persist(task)
        log(f"Saved: {output_path}")
        return output_path
    except Exception as exc:
        task["error"] = str(exc)[:2000]
        if not task.get("provider_task_id"):
            task["status"] = "submission_unknown"
        elif str(task.get("provider_status", "")).casefold() in {
            "failed", "error", "cancelled", "canceled"
        }:
            task["status"] = "failed"
        else:
            task["status"] = "polling_interrupted"
        persist(task)
        raise


def run_command(args: argparse.Namespace) -> int:
    inventory_path = Path(args.inventory).resolve()
    selections_path = Path(args.selections).resolve()
    inventory = json_load(inventory_path)
    selections = json_load(selections_path)
    errors = validate_selection_data(inventory, selections)
    if errors:
        for error in errors:
            log(f"ERROR: {error}")
        raise WorkflowError("选择清单校验失败；未调用任何付费 API")

    work_dir = Path(inventory["work_dir"])
    with WorkDirectoryLock(work_dir / ".concurrent_run.lock"):
        return run_command_locked(args, inventory, selections, work_dir)


def run_command_locked(
    args: argparse.Namespace,
    inventory: dict[str, Any],
    selections: dict[str, Any],
    work_dir: Path,
) -> int:
    checkpoint_path = work_dir / "checkpoint.json"
    if checkpoint_path.is_file():
        checkpoint = json_load(checkpoint_path)
    else:
        checkpoint = {
            "schema_version": 1,
            "run_id": inventory["run_id"],
            "config": {
                "language_model": args.language_model,
                "generation_model": args.generation_model,
                "aspect_ratio": args.aspect_ratio,
                "resolution": args.resolution,
            },
            "upload_cache": {},
            "tasks": {},
        }
        checkpoint_save(checkpoint_path, checkpoint)

    store = CheckpointStore(checkpoint_path, checkpoint)
    config = checkpoint["config"]
    if args.resume:
        language_model = config["language_model"]
        generation_model = config["generation_model"]
        aspect_ratio = config["aspect_ratio"]
        resolution = config["resolution"]
        log(f"Resume uses checkpoint config: {config}")
    else:
        requested = {
            "language_model": args.language_model,
            "generation_model": args.generation_model,
            "aspect_ratio": args.aspect_ratio,
            "resolution": args.resolution,
        }
        if config != requested and checkpoint.get("tasks"):
            raise WorkflowError("已有 checkpoint 任务且配置不同；请使用 --resume，勿重复提交")
        pending = [
            task for task in checkpoint.get("tasks", {}).values()
            if task.get("provider_task_id")
            and not (
                task.get("status") == "completed"
                and Path(task.get("output_path", "")).is_file()
            )
        ]
        if pending:
            raise WorkflowError("检测到未完成 provider task；请使用 --resume，勿重复提交")
        checkpoint["config"] = requested
        language_model, generation_model, aspect_ratio, resolution = (
            args.language_model, args.generation_model, args.aspect_ratio, args.resolution
        )
        store.save()

    checkpoint["concurrency"] = {
        "engine": "thread_pool",
        "max_workers": args.max_workers,
    }
    store.save()

    toapis_key = prompt_secret("ToAPIs API Key", "TOAPIS_API_KEY")
    runninghub_key = ""
    if generation_model == "gpt-image-2-rh":
        runninghub_key = prompt_secret("RunningHub API Key", "RUNNINGHUB_API_KEY")

    inventory_by_sku = {entry["sku"]: entry for entry in inventory["skus"]}
    completed = 0
    jobs: list[dict[str, Any]] = []
    for sku_selection in selections["skus"]:
        sku = sku_selection["sku"]
        sku_dir = Path(inventory_by_sku[sku]["folder"])
        extra_requirements = read_sku_prompt(sku_dir)
        for model in sku_selection["models"]:
            key = task_key(sku, model["path"])
            output_path = sku_dir / f"result_{inventory['run_id']}_{Path(model['path']).stem}.png"
            task = store.get_task(key)
            if task and task.get("status") == "completed" and output_path.is_file():
                log(f"Skip completed: {sku}/{Path(model['path']).name}")
                completed += 1
                continue
            if task and task.get("provider_task_id"):
                if not args.resume:
                    raise WorkflowError(f"{sku}/{Path(model['path']).name} already submitted; use --resume")
                jobs.append({"key": key, "output_path": str(output_path)})
                continue
            if task and task.get("status") in {"submission_unknown", "failed"}:
                raise WorkflowError(
                    f"{sku}/{Path(model['path']).name} 状态为 {task['status']} 且没有可恢复任务 ID；"
                    "为防止重复扣费，不自动重提"
                )
            if task and task.get("status") == "prepared":
                prompt_path = Path(task.get("prompt_path", ""))
                if task.get("ordered_image_urls") and prompt_path.is_file():
                    task["prompt"] = prompt_path.read_text(encoding="utf-8")
                    figure_plan_path = Path(task.get("figure_plan_path", ""))
                    if figure_plan_path.is_file():
                        prepared_plan = json_load(figure_plan_path)
                        contract_errors = validate_prompt_contracts(task["prompt"], prepared_plan)
                        if contract_errors:
                            raise WorkflowError(
                                "已准备任务的提示词裤脚约束校验失败；未提交付费生图: "
                                + "; ".join(contract_errors)
                            )
                    store.set_task(key, task)
                    jobs.append({"key": key, "output_path": str(output_path)})
                    continue

            figure_plan = build_figure_plan(model, sku_selection["product_groups"])
            coverage_errors = validate_figure_product_coverage(
                sku, sku_selection["product_groups"], figure_plan
            )
            if coverage_errors:
                raise WorkflowError(
                    "产品 Figure 完整性校验失败；未上传图片且未调用付费 API: "
                    + "; ".join(coverage_errors)
                )
            figure_plan_path = work_dir / "figure_plans" / sku / f"{Path(model['path']).stem}.json"
            json_save(figure_plan_path, figure_plan)
            log(f"Figure plan ready: {figure_plan_path}")
            for item in figure_plan:
                if item["role"] == "fit_reference":
                    log(f"  Figure {item['figure']}: fit reference only for Figure {item['for_product_figure']} — {Path(item['path']).name}")
                else:
                    log(f"  Figure {item['figure']}: {item['role']} — {Path(item['path']).name}")
                    if item.get("lower_hem"):
                        log(
                            f"    reviewed lower_hem={item['lower_hem']['type']} — "
                            f"{item['lower_hem']['reason']}"
                        )

            image_urls: list[str] = []
            for item in figure_plan:
                image_path = Path(item["path"])
                cache_key = path_key(image_path)
                url = checkpoint["upload_cache"].get(cache_key)
                if not url:
                    log(f"Uploading Figure {item['figure']}: {image_path.name}")
                    url = upload_image(toapis_key, image_path)
                    checkpoint["upload_cache"][cache_key] = url
                    store.save()
                image_urls.append(url)

            mapping_text = build_mapping_text(figure_plan, extra_requirements)
            prompt = generate_prompt(toapis_key, language_model, image_urls, mapping_text)
            prompt = apply_lower_hem_contracts(prompt, figure_plan)
            contract_errors = validate_prompt_contracts(prompt, figure_plan)
            if contract_errors:
                raise WorkflowError(
                    "最终提示词裤脚约束校验失败；未提交付费生图: "
                    + "; ".join(contract_errors)
                )
            prompt_path = work_dir / "prompts" / sku / f"{Path(model['path']).stem}.txt"
            prompt_path.parent.mkdir(parents=True, exist_ok=True)
            prompt_path.write_text(prompt, encoding="utf-8")
            log(f"Prompt ready: {prompt_path}")

            task = {
                "sku": sku,
                "model_path": str(Path(model["path"]).resolve()),
                "output_path": str(output_path),
                "figure_plan_path": str(figure_plan_path),
                "prompt_path": str(prompt_path),
                "ordered_image_urls": image_urls,
                "provider": "runninghub" if generation_model == "gpt-image-2-rh" else "toapis",
                "status": "prepared",
                "prompt": prompt,
            }
            store.set_task(key, task)
            jobs.append({"key": key, "output_path": str(output_path)})

    if not jobs:
        log(f"Run complete: {completed} model result(s); nothing pending")
        return 0

    actual_workers = min(args.max_workers, len(jobs))
    log(
        f"Starting concurrent paid stage: jobs={len(jobs)}, "
        f"max_workers={actual_workers}, provider={generation_model}"
    )
    failures: list[tuple[str, str]] = []
    with ThreadPoolExecutor(
        max_workers=actual_workers,
        thread_name_prefix="fit-tryon",
    ) as executor:
        future_to_job = {
            executor.submit(
                execute_generation_job,
                job,
                store,
                toapis_key,
                runninghub_key,
                generation_model,
                aspect_ratio,
                resolution,
            ): job
            for job in jobs
        }
        for future in as_completed(future_to_job):
            job = future_to_job[future]
            try:
                future.result()
                completed += 1
            except Exception as exc:
                failures.append((job["key"], str(exc)))
                log(f"ERROR worker {job['key']}: {exc}")

    if failures:
        details = "; ".join(f"{key}: {message}" for key, message in failures[:5])
        raise WorkflowError(
            f"并发阶段完成，但 {len(failures)} 个任务失败；其他任务已保留。{details}"
        )

    log(
        f"Run complete: {completed} model result(s); "
        f"concurrent workers={actual_workers}"
    )
    return 0


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    subparsers = parser.add_subparsers(dest="command", required=True)

    prepare_parser = subparsers.add_parser("prepare", help="Inventory SKU folders and build contact sheets")
    prepare_parser.add_argument("--root", required=True, help="Target root containing immediate SKU child folders")
    table_group = prepare_parser.add_mutually_exclusive_group(required=True)
    table_group.add_argument("--table-excel", help="Local 亚马逊视觉需求表 workbook")
    table_group.add_argument("--table-json", help="Normalized DingTalk read-only JSON")
    prepare_parser.add_argument("--library", required=True, help="Fit-reference library root")
    prepare_parser.set_defaults(func=prepare)

    validate_parser = subparsers.add_parser("validate-selections", help="Validate reviewed selections before paid API calls")
    validate_parser.add_argument("--inventory", required=True)
    validate_parser.add_argument("--selections", required=True)
    validate_parser.set_defaults(func=validate_command)

    run_parser = subparsers.add_parser("run", help="Generate prompts and images from a validated selection manifest")
    run_parser.add_argument("--inventory", required=True)
    run_parser.add_argument("--selections", required=True)
    run_parser.add_argument("--language-model", default="gemini-3.1-flash-lite", choices=["gemini-3.1-flash-lite", "gpt-5.6-sol"])
    run_parser.add_argument("--generation-model", default="gemini-3.1-flash-image-preview", choices=["gemini-3.1-flash-image-preview", "gpt-image-2", "gpt-image-2-rh"])
    run_parser.add_argument("--aspect-ratio", default="3:4")
    run_parser.add_argument("--resolution", default="4K", choices=["1K", "2K", "4K"])
    run_parser.add_argument(
        "--max-workers",
        type=worker_count,
        default=3,
        help="Concurrent paid generation workers (1-8, default: 3)",
    )
    run_parser.add_argument("--resume", action="store_true", help="Resume existing submitted tasks; never resubmit them")
    run_parser.set_defaults(func=run_command)
    return parser


def main() -> int:
    parser = build_parser()
    args = parser.parse_args()
    try:
        return int(args.func(args))
    except KeyboardInterrupt:
        log("Interrupted. Check checkpoint.json and resume instead of resubmitting.")
        return 130
    except WorkflowError as exc:
        log(f"ERROR: {exc}")
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
